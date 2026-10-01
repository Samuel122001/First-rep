#!/usr/bin/env python3
"""Builds initial PMI Hub data from an employee list and a PMI Gantt workbook.

Usage:
    python3 scripts/excel_to_seed.py <gantt.xlsx> <employees.txt> <out.json> [mapping.json]

Real company data must stay out of git: keep the inputs and the output in
seed/ under the *.local.* / local-seed.json names (see seed/README.md).

The workbook is expected to have:
  B2/C2  START DATE        B3/C3  PROJECT TITLE      B4/C4  PROJECT MANAGER
  row 6  headers: TASK TITLE | TASK DESCRIPTION | DEFINITION OF DONE | TASK OWNER |
         SCHEDULED START | SCHEDULED FINISH | OPEN / DONE
  Phase rows:       a number in column A ("1.", "2."), the phase name in column B
  Workstream rows:  "Name -- Lead" in column B, no dates
  Task rows:        any other row with text in column B

Owners in the workbook are usually first names ("Anna / Erik"). Each is matched
to exactly one person on the employee list (first name first, then any part of
the name, ignoring accents). Owners without a unique match are added as
separate people and reported, so no assignment from the workbook is lost.

mapping.json (optional):
  {"projectId": "p_acme", "company": "Acme", "source": "Acme_PMI.xlsx",
   "aliases": {"Misspelt": "Name"},
   "notOnList": {"Law firm X": {"type": "external", "role": "Law firm", "team": ""}},
   "milestones": ["SPA Signing"]}

The output is a list of documents {path, data} for the app's database
(people/, projects/, workstreams/, tasks/, history/).
"""
import json
import re
import sys
import unicodedata
from datetime import datetime, timezone

import openpyxl

IMPORT_TS = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")
CFG = {"projectId": "p_project", "company": "", "source": "workbook.xlsx", "aliases": {}, "notOnList": {}, "milestones": ["SPA Signing"]}
# Workstreams in display order: workbook name -> app name.
WORKSTREAMS = [
    ("M&A Team & IC", "M&A & IC"),
    ("Communication", "Communication"),
    ("Commercial", "Commercial"),
    ("Operations", "Operations"),
    ("Finance", "Finance"),
    ("People", "People"),
    ("Tech", "Tech"),
    ("Analytics", "Analytics"),
]


def fold(s):
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c)).lower()


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", fold(s).replace("&", "and")).strip("-")


def iso(v):
    return v.strftime("%Y-%m-%d") if isinstance(v, datetime) else None


def split_owners(text):
    if not text:
        return []
    text = text.strip()
    if text == "M&A & IC":
        return ["M&A Team", "IC"]
    return [CFG["aliases"].get(p.strip(), p.strip()) for p in text.split("/") if p.strip()]


def match_employee(short, employees):
    key = fold(short)
    first = [e for e in employees if fold(e.split()[0]) == key]
    if len(first) == 1:
        return first[0]
    anywhere = [e for e in employees if key in [fold(p) for p in e.split()]]
    return anywhere[0] if len(anywhere) == 1 else None


def person_doc(name, ptype="employee", role="", team="", source="list"):
    return {
        "name": name,
        "role": role,
        "team": team,
        "email": "",
        "type": ptype,
        "active": True,
        "userId": None,
        "source": source,
        "createdAt": IMPORT_TS,
        "createdBy": None,
    }


def main(xlsx, employees_file, dst, mapping_file=None):
    if mapping_file:
        CFG.update(json.load(open(mapping_file, encoding="utf-8")))
    PROJECT_ID, SOURCE, NOT_ON_LIST, MILESTONES = CFG["projectId"], CFG["source"], CFG["notOnList"], set(CFG["milestones"])
    key = PROJECT_ID.removeprefix("p_")
    employees = [l.strip() for l in open(employees_file, encoding="utf-8") if l.strip()]
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.worksheets[0]
    project_title = ws["C3"].value or "PMI"
    project_manager = ws["C4"].value or ""

    phases, ws_leads, tasks = [], {}, []
    cur_phase = cur_ws = None
    order = 0
    for row in range(8, ws.max_row + 1):
        a = ws.cell(row, 1).value
        b = ws.cell(row, 2).value
        b = b.strip() if isinstance(b, str) else b
        if not b:
            continue
        if a and re.match(r"^\d+\.?$", str(a).strip()):
            name = b.replace("Pre - Acquisition", "Pre-Acquisition")
            cur_phase = {"id": "ph_" + slug(name), "name": name}
            phases.append(cur_phase)
            continue
        if " -- " in b:
            ws_name, lead = [x.strip() for x in b.split(" -- ", 1)]
            cur_ws = dict(WORKSTREAMS).get(ws_name, ws_name)
            ws_leads.setdefault(cur_ws, lead)
            continue
        c, d, e, f, g, h = (ws.cell(row, col).value for col in range(3, 9))
        order += 1
        tasks.append({
            "row": row, "phase": cur_phase["id"], "ws": cur_ws, "title": b,
            "description": (c or "").strip(), "dod": (d or "").strip(), "owners": split_owners(e),
            "start": iso(f), "due": iso(g),
            "status": "done" if (h or "").strip().lower() == "done" else "open",
            "order": order * 10,
        })

    # Everyone on the employee list.
    people = {}  # person id -> doc
    for name in employees:
        people["pe_" + slug(name)] = person_doc(name)

    # Map every short name used in the workbook to a person id.
    short_names = []
    for t in tasks:
        short_names += t["owners"]
    short_names += list(ws_leads.values()) + ([project_manager] if project_manager else [])
    resolved, report = {}, []
    for short in dict.fromkeys(short_names):
        full = None if short in NOT_ON_LIST else match_employee(short, employees)
        if full:
            resolved[short] = "pe_" + slug(full)
            report.append(f"  {short:<15} -> {full}")
        else:
            extra = NOT_ON_LIST.get(short, {"type": "employee", "role": "", "team": ""})
            pid = "pe_" + slug(short)
            people[pid] = person_doc(short, extra["type"], extra["role"], extra["team"], source="excel")
            resolved[short] = pid
            report.append(f"  {short:<15} -> (not on the list, added: {extra['type']})")

    docs = []
    for pid, body in people.items():
        docs.append({"path": f"people/{pid}", "data": body})
        docs.append({"path": f"history/{pid}", "data": {
            "projectId": "_people", "type": "person", "entityId": pid,
            "events": {"e0import": {"t": IMPORT_TS, "u": None, "a": "import", "l": body["name"],
                                    "src": "employee list" if body["source"] == "list" else SOURCE}},
        }})

    plan_start = min((t["start"] for t in tasks if t["start"]), default=None)
    plan_end = max((t["due"] for t in tasks if t["due"]), default=None)
    closing = next((t["due"] for t in tasks if t["title"] in MILESTONES and t["due"]), None)

    docs.append({"path": f"projects/{PROJECT_ID}", "data": {
        "name": project_title, "company": CFG["company"], "status": "active",
        "signingDate": closing, "closingDate": closing, "planStart": plan_start, "planEnd": plan_end,
        "leadPersonId": resolved.get(project_manager), "description": "", "phases": phases,
        "createdAt": IMPORT_TS, "createdBy": None, "updatedAt": IMPORT_TS, "updatedBy": None, "source": SOURCE,
    }})
    docs.append({"path": f"history/{PROJECT_ID}", "data": {
        "projectId": PROJECT_ID, "type": "project", "entityId": PROJECT_ID,
        "events": {"e0import": {"t": IMPORT_TS, "u": None, "a": "import", "l": project_title, "src": SOURCE, "n": len(tasks)}},
    }})

    ws_ids = {}
    for i, (_, app_name) in enumerate(WORKSTREAMS):
        wid = f"w_{key}_" + slug(app_name)
        ws_ids[app_name] = wid
        lead = ws_leads.get(app_name)
        docs.append({"path": f"workstreams/{wid}", "data": {
            "projectId": PROJECT_ID, "name": app_name, "leadPersonId": resolved.get(lead), "description": "",
            "order": (i + 1) * 10, "rag": None, "ragNote": "", "ragAt": None, "ragBy": None, "deleted": False,
            "createdAt": IMPORT_TS, "createdBy": None,
        }})
        docs.append({"path": f"history/{wid}", "data": {
            "projectId": PROJECT_ID, "type": "workstream", "entityId": wid,
            "events": {"e0import": {"t": IMPORT_TS, "u": None, "a": "import", "l": app_name, "src": SOURCE}},
        }})

    for t in tasks:
        tid = f"t_{key}_r{t['row']}"
        docs.append({"path": f"tasks/{tid}", "data": {
            "projectId": PROJECT_ID, "workstreamId": ws_ids[t["ws"]], "phaseId": t["phase"],
            "title": t["title"], "description": t["description"], "dod": t["dod"],
            "ownerIds": list(dict.fromkeys(resolved[o] for o in t["owners"])),
            "start": t["start"], "due": t["due"], "status": t["status"], "priority": "normal",
            "milestone": t["title"] in MILESTONES, "dependsOn": [], "order": t["order"], "deleted": False,
            "doneAt": None, "createdAt": IMPORT_TS, "createdBy": None, "updatedAt": IMPORT_TS, "updatedBy": None,
        }})
        docs.append({"path": f"history/{tid}", "data": {
            "projectId": PROJECT_ID, "type": "task", "entityId": tid,
            "events": {"e0import": {"t": IMPORT_TS, "u": None, "a": "import", "l": t["title"], "src": SOURCE, "row": t["row"]}},
        }})

    with open(dst, "w", encoding="utf-8") as fh:
        json.dump({"generatedAt": datetime.now(timezone.utc).isoformat(), "source": SOURCE, "docs": docs}, fh, ensure_ascii=False, indent=1)
    kinds = {}
    for d in docs:
        k = d["path"].split("/")[0]
        kinds[k] = kinds.get(k, 0) + 1
    print(f"{len(docs)} documents {kinds}")
    print(f"{len(employees)} people from the list, {len(people) - len(employees)} added from the workbook, {len(tasks)} tasks")
    print("Workbook owners:")
    print("\n".join(report))


if __name__ == "__main__":
    main(*sys.argv[1:5])
